#!/usr/bin/env python3
"""Tests du classeur d'annonces (commun/classeur.py et classeur_ads.py) :
python3 -m unittest discover evals/chatgpt-ads-excel

Classeurs abîmés construits à la main (XML brut, sans dépendance) : italique,
fusion, formule, identifiant numérique, onglet en trop, en-têtes renommés ou
déplacés, CSV ; puis réparation, écriture, dimensions d'images et contrôle
en ligne simulé.
"""
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

PACK = Path(__file__).resolve().parents[2]
SKILL = PACK / "skills" / "chatgpt-ads-excel"
SCRIPT = SKILL / "scripts" / "classeur_ads.py"
sys.path.insert(0, str(SKILL / "commun"))
import classeur as cl  # noqa: E402

BON = {"ad_id": "", "ad_group_id": "adgrp_101", "adgroup_name": "expat-preparer", "ad_name": "expat-preparer-01",
       "status": "active", "title": "Partir vivre à Lisbonne", "copy": "Les démarches d'installation, étape par étape.",
       "link": "https://assurly.example/expatriation?utm_source=chatgpt", "image_url": "https://assurly.example/img/expat.png",
       "image_file_id": ""}


def col(n):
    return cl._col_lettre(n)


def classeur_brut(chemin, lignes, entetes=None, onglets_en_plus=(), italique=(), fusion=None, formule=None,
                  nombres=(), nom="Feuil1"):
    """XLSX écrit à la main. italique/nombres : {(ligne, colonne)} ; formule : (ligne, colonne)."""
    entetes = entetes or cl.COLONNES
    def cel(r, c, v):
        ref = f"{col(c)}{r}"
        s = ' s="1"' if (r, c) in italique else ""
        if formule == (r, c):
            return f'<c r="{ref}"{s} t="str"><f>CONCAT("a","b")</f><v>{escape(v)}</v></c>'
        if (r, c) in nombres:
            return f'<c r="{ref}"{s}><v>{v}</v></c>'
        return f'<c r="{ref}"{s} t="inlineStr"><is><t>{escape(v)}</t></is></c>'
    rangs = ["<row r=\"1\">" + "".join(cel(1, j, h) for j, h in enumerate(entetes, 1)) + "</row>"]
    for i, l in enumerate(lignes, 2):
        rangs.append(f'<row r="{i}">' + "".join(cel(i, j, l.get(h, "")) for j, h in enumerate(entetes, 1) if l.get(h, "") != "") + "</row>")
    fus = f'<mergeCells count="1"><mergeCell ref="{fusion}"/></mergeCells>' if fusion else ""
    feuille = ('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
               + "".join(rangs) + "</sheetData>" + fus + "</worksheet>")
    noms = [nom] + list(onglets_en_plus)
    wb = ('<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets>'
          + "".join(f'<sheet name="{n}" sheetId="{i}" r:id="rId{i}"/>' for i, n in enumerate(noms, 1)) + "</sheets></workbook>")
    rels = ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            + "".join(f'<Relationship Id="rId{i}" Type="t" Target="worksheets/sheet{i}.xml"/>' for i in range(1, len(noms) + 1))
            + "</Relationships>")
    styles = ('<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
              '<fonts count="2"><font><sz val="11"/></font><font><i/><sz val="11"/></font></fonts>'
              '<cellXfs count="2"><xf fontId="0"/><xf fontId="1"/></cellXfs></styleSheet>')
    with zipfile.ZipFile(chemin, "w") as z:
        z.writestr("xl/workbook.xml", wb)
        z.writestr("xl/_rels/workbook.xml.rels", rels)
        z.writestr("xl/styles.xml", styles)
        z.writestr("xl/worksheets/sheet1.xml", feuille)
        for i in range(2, len(noms) + 1):
            z.writestr(f"xl/worksheets/sheet{i}.xml", '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData/></worksheet>')


def messages(r, niveau=None):
    return [p["message"] for p in r.points if niveau in (None, p["niveau"])]


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def verif(self, lignes, **kw):
        f = self.d / "c.xlsx"
        classeur_brut(f, lignes, **kw)
        r, _, l = cl.verifier(f)
        return r, l


class TestStructure(Base):
    def test_classeur_propre(self):
        r, l = self.verif([BON])
        self.assertEqual(r.compte(cl.BLOQUANT), 0, r.texte())
        self.assertEqual(r.compte(cl.ATTENTION), 0, r.texte())
        self.assertEqual(l[0]["title"], BON["title"])

    def test_onglet_en_trop(self):
        r, _ = self.verif([BON], onglets_en_plus=["notes"])
        self.assertTrue(any("onglets en trop : notes" in m for m in messages(r, cl.BLOQUANT)))

    def test_entete_renommee(self):
        e = list(cl.COLONNES)
        e[2] = "ad_group_name"
        r, _ = self.verif([], entetes=e)
        self.assertTrue(any("ad_group_name" in m and "adgroup_name" in m for m in messages(r, cl.BLOQUANT)), r.texte())

    def test_entete_proche(self):
        e = list(cl.COLONNES)
        e[5] = "Title"
        r, _ = self.verif([], entetes=e)
        self.assertTrue(any("« Title » au lieu de « title »" in m for m in messages(r, cl.BLOQUANT)), r.texte())

    def test_ordre(self):
        e = list(cl.COLONNES)
        e[5], e[6] = e[6], e[5]
        r, _ = self.verif([BON], entetes=e)
        self.assertTrue(any("autre ordre" in m for m in messages(r, cl.BLOQUANT)))

    def test_colonne_en_trop(self):
        r, _ = self.verif([BON], entetes=cl.COLONNES + ["notes"])
        self.assertTrue(any("« notes » hors schéma" in m for m in messages(r, cl.BLOQUANT)))

    def test_fusion(self):
        r, _ = self.verif([BON], fusion="F2:G2")
        self.assertTrue(any("fusionnées" in m for m in messages(r, cl.BLOQUANT)))

    def test_aucune_annonce(self):
        r, _ = self.verif([])
        self.assertIn("aucune annonce", messages(r, cl.BLOQUANT))
        self.assertIn("classeur XLSX : le format d'import retenu est le CSV", messages(r, cl.INFO))


class TestCellules(Base):
    def test_italique(self):
        r, _ = self.verif([BON], italique={(2, 6)})
        self.assertTrue(any("title : en italique" in m for m in messages(r, cl.ATTENTION)))

    def test_formule(self):
        r, _ = self.verif([dict(BON, ad_name="ab")], formule=(2, 4))
        self.assertTrue(any("ad_name : formule" in m for m in messages(r, cl.BLOQUANT)))

    def test_identifiant_nombre(self):
        r, l = self.verif([dict(BON, ad_group_id="1234567890123")], nombres={(2, 2)})
        self.assertTrue(any("ad_group_id : enregistré comme nombre" in m for m in messages(r, cl.ATTENTION)))
        self.assertEqual(cl.texte(l[0]["ad_group_id"]), "1234567890123")

    def test_espaces_et_invisibles(self):
        r, _ = self.verif([dict(BON, title=" Partir vivre​ ")])
        m = messages(r, cl.ATTENTION)
        self.assertTrue(any("espace au début" in x for x in m))
        self.assertTrue(any("invisible" in x for x in m))

    def test_retour_ligne(self):
        r, _ = self.verif([dict(BON, copy="Ligne un\nligne deux")])
        self.assertTrue(any("retour à la ligne" in m for m in messages(r, cl.BLOQUANT)))


class TestAnnonces(Base):
    def test_longueurs(self):
        r, _ = self.verif([dict(BON, title="Un titre beaucoup trop long ici", copy="x" * 49)])
        b = messages(r, cl.BLOQUANT)
        self.assertTrue(any(m.startswith("title : 31 caractères pour 24") for m in b))
        self.assertTrue(any(m.startswith("copy : 49 caractères pour 48") for m in b))

    def test_limites_api(self):
        f = self.d / "a.xlsx"
        classeur_brut(f, [dict(BON, title="Un titre beaucoup trop long ici")])
        r, _, _ = cl.verifier(f, "api")
        self.assertEqual(r.compte(cl.BLOQUANT), 0, r.texte())

    def test_accents_comptes_une_fois(self):
        import unicodedata
        t = unicodedata.normalize("NFD", "Été à Évora, très doré")   # 22 caractères affichés
        self.assertEqual(cl.longueur(t), 22)

    def test_typographie_francaise(self):
        self.assertEqual(cl.longueur("Pourquoi attendre ?"), 19)

    def test_groupe_deux_noms(self):
        r, _ = self.verif([BON, dict(BON, adgroup_name="Expat-preparer", ad_name="x-02", title="Autre titre ici")])
        self.assertTrue(any("porte deux noms" in m for m in messages(r, cl.BLOQUANT)))

    def test_nom_deux_identifiants(self):
        r, _ = self.verif([BON, dict(BON, ad_group_id="adgrp_102", ad_name="x-02", title="Autre titre ici")])
        self.assertTrue(any("relié à deux identifiants" in m for m in messages(r, cl.BLOQUANT)))

    def test_cases_vides_listees_sans_bloquer(self):
        r, _ = self.verif([dict(BON, ad_group_id="", adgroup_name="", link="", image_url=""),
                           dict(BON, ad_name="b-02", title="Autre titre ici", image_url="", image_file_id="file_3")])
        self.assertEqual(r.compte(cl.BLOQUANT), 0, r.texte())
        a = messages(r, cl.ATTENTION)
        self.assertIn("adgroup_name vide sur 1 ligne(s) : 2", a)
        self.assertIn("link vide sur 1 ligne(s) : 2", a)
        self.assertIn("image_url vide sur 1 ligne(s) : 2", a)       # la ligne 3 a un image_file_id

    def test_sans_identifiant_de_groupe(self):
        r, _ = self.verif([dict(BON, ad_group_id="")])
        self.assertTrue(any(m.startswith("ad_group_id vide sur 1") for m in messages(r, cl.INFO)))
        self.assertEqual(r.compte(cl.ATTENTION), 0, r.texte())

    def test_prefixe(self):
        r, _ = self.verif([dict(BON, ad_group_id="101", image_file_id="img9")])
        a = messages(r, cl.ATTENTION)
        self.assertTrue(any("« adgrp_ »" in m for m in a))
        self.assertTrue(any("« file_ »" in m for m in a))

    def test_statut(self):
        r, _ = self.verif([dict(BON, status="en pause")])
        self.assertTrue(any("status « en pause »" in m for m in messages(r, cl.BLOQUANT)))
        r, _ = self.verif([dict(BON, status="Paused")])
        self.assertTrue(any("majuscules" in m for m in messages(r, cl.ATTENTION)))
        r, _ = self.verif([dict(BON, status="paused")])
        self.assertTrue(any("status paused sur 1" in m for m in messages(r, cl.INFO)))
        r, _ = self.verif([dict(BON, status="active")])
        self.assertFalse(any("status" in m for m in messages(r)))

    def test_liens(self):
        r, _ = self.verif([dict(BON, link="assurly.example/expat")])
        self.assertTrue(any("link « assurly.example/expat »" in m for m in messages(r, cl.BLOQUANT)))
        r, _ = self.verif([dict(BON, link="http://assurly.example/expat")])
        self.assertTrue(any("link en http" in m for m in messages(r, cl.ATTENTION)))

    def test_image(self):
        r, _ = self.verif([dict(BON, image_url="")])
        self.assertIn("image_url vide sur 1 ligne(s) : 2", messages(r, cl.ATTENTION))
        r, _ = self.verif([dict(BON, image_url="", image_file_id="file_9")])
        self.assertEqual(r.compte(cl.BLOQUANT), 0, r.texte())
        r, _ = self.verif([dict(BON, image_url="https://drive.example/partage?id=9")])
        self.assertTrue(any("sans extension" in m for m in messages(r, cl.ATTENTION)))

    def test_doublons(self):
        r, _ = self.verif([BON, BON])
        b = messages(r, cl.BLOQUANT)
        self.assertTrue(any("double exact" in m for m in b))
        r, _ = self.verif([BON, dict(BON, ad_name="expat-preparer-02")])
        self.assertTrue(any("même titre et même copie" in m for m in messages(r, cl.ATTENTION)))

    def test_ad_id_en_double(self):
        r, _ = self.verif([dict(BON, ad_id="ad_1"), dict(BON, ad_id="ad_1", ad_name="b-02", title="Autre titre")])
        self.assertTrue(any("ad_id ad_1 déjà" in m for m in messages(r, cl.BLOQUANT)))

    def test_csv_point_virgule(self):
        f = self.d / "c.csv"
        f.write_text(";".join(cl.COLONNES) + "\n" + ";".join(BON[c] for c in cl.COLONNES) + "\n", encoding="utf-8")
        r, _, l = cl.verifier(f)
        self.assertEqual(r.compte(cl.BLOQUANT), 0, r.texte())
        self.assertEqual(l[0]["link"], BON["link"])


class TestCSV(Base):
    def test_ecriture_csv(self):
        f = self.d / "i.csv"
        cl.ecrire(f, [dict(BON, copy="Démarches, budget et calendrier.")])
        brut = f.read_bytes()
        self.assertFalse(brut.startswith(b"\xef\xbb\xbf"))
        self.assertTrue(brut.decode("utf-8").startswith(",".join(cl.COLONNES) + "\r\n"))
        r, classeur, l = cl.verifier(f)
        self.assertEqual(r.compte(cl.BLOQUANT) + r.compte(cl.ATTENTION), 0, r.texte())
        self.assertEqual(l[0]["copy"], "Démarches, budget et calendrier.")

    def test_csv_point_virgule_bom_signales(self):
        f = self.d / "c.csv"
        f.write_bytes(b"\xef\xbb\xbf" + (";".join(cl.COLONNES) + "\n" + ";".join(BON[c] for c in cl.COLONNES) + "\n").encode())
        r, _, l = cl.verifier(f)
        a = messages(r, cl.ATTENTION)
        self.assertTrue(any("séparateur « ; »" in m for m in a))
        self.assertTrue(any("BOM" in m for m in a))
        self.assertEqual(l[0]["ad_id"], "")

    def test_reparer_xlsx_vers_csv(self):
        f, s = self.d / "c.xlsx", self.d / "c.csv"
        classeur_brut(f, [dict(BON, status="")])
        p = subprocess.run([sys.executable, str(SCRIPT), "reparer", str(f), "-o", str(s)], capture_output=True, text=True)
        self.assertEqual(p.returncode, 0, p.stdout)
        _, l = cl.annonces(cl.lire(s))
        self.assertEqual((l[0]["status"], l[0]["image_file_id"], l[0]["ad_id"]), ("active", "", ""))


class TestReparation(Base):
    def test_reparer_puis_relire(self):
        f, s = self.d / "sale.xlsx", self.d / "propre.xlsx"
        lignes = [dict(BON, status="", ad_name="", title=" Partir vivre à Lisbonne ", copy="Démarches, budget et calendrier du départ."), dict(BON, status="Paused"),
                  dict(BON, status="Paused")]
        classeur_brut(f, lignes, onglets_en_plus=["brouillon"], italique={(2, 6)}, nombres=set())
        p = subprocess.run([sys.executable, str(SCRIPT), "reparer", str(f), "-o", str(s)], capture_output=True, text=True)
        self.assertIn("status vide → active", p.stdout)
        self.assertNotIn("déjà utilisé", p.stdout)
        self.assertIn("doublon exact retiré", p.stdout)
        r, cl_, l = cl.verifier(s)
        self.assertEqual(r.compte(cl.BLOQUANT), 0, r.texte())
        self.assertEqual([o["nom"] for o in cl_["onglets"]], ["Feuil1"])
        self.assertEqual(len(l), 2)
        self.assertEqual(l[0]["title"], "Partir vivre à Lisbonne")
        self.assertTrue(all(c["type"] == "texte" and not c["italique"] for x in l for c in x["_cellules"].values()))

    def test_reparer_ne_coupe_pas(self):
        long = "Un titre beaucoup trop long ici"
        propres, _ = cl.reparer_lignes([dict(BON, title=long)])
        self.assertEqual(propres[0]["title"], long)

    def test_reparer_idempotent(self):
        p1, _ = cl.reparer_lignes([dict(BON, status="", ad_name="")])
        p2, j2 = cl.reparer_lignes(p1)
        self.assertEqual(p1, p2)
        self.assertEqual(j2, [])

    def test_ecriture_metadonnees(self):
        f = self.d / "e.xlsx"
        cl.ecrire_xlsx(f, [BON])
        with zipfile.ZipFile(f) as z:
            self.assertNotIn("docProps/core.xml", z.namelist())
            self.assertNotIn("<i/>", z.read("xl/styles.xml").decode())
        try:
            import openpyxl
        except ImportError:
            return
        ws = openpyxl.load_workbook(f).active
        self.assertEqual([c.value for c in ws[1]], cl.COLONNES)
        self.assertEqual(ws["F2"].value, BON["title"])

    def test_echappement(self):
        f = self.d / "e.xlsx"
        cl.ecrire_xlsx(f, [dict(BON, title="Prix & délais <clairs>")])
        _, l = cl.annonces(cl.lire(f))
        self.assertEqual(l[0]["title"], "Prix & délais <clairs>")


class TestCommandes(Base):
    def run_(self, *args):
        return subprocess.run([sys.executable, str(SCRIPT), *map(str, args)], capture_output=True, text=True)

    def test_schema(self):
        f = self.d / "v.xlsx"
        p = self.run_("schema", "-o", f)
        self.assertEqual(p.returncode, 0)
        r, _, _ = cl.verifier(f)
        self.assertIn("aucune annonce", messages(r))

    def test_verifier_codes(self):
        f = self.d / "c.xlsx"
        classeur_brut(f, [BON])
        self.assertEqual(self.run_("verifier", f).returncode, 0)
        classeur_brut(f, [dict(BON, title="x" * 30)])
        p = self.run_("verifier", f, "--json")
        self.assertEqual(p.returncode, 3)
        self.assertEqual(json.loads(p.stdout)["bloquants"], 1)
        self.assertEqual(self.run_("verifier", self.d / "absent.xlsx").returncode, 1)

    def test_depuis_json(self):
        camp = {"adgroups": [{"adgroup_name": "expat-preparer", "ad_group_id": "adgrp_101"}],
                "ads": [{"adgroup_name": "expat-preparer", "ad_name": "expat-preparer-01", "title": BON["title"],
                         "copy": BON["copy"], "link": BON["link"], "image_link": BON["image_url"]}]}
        j, s = self.d / "c.json", self.d / "s.xlsx"
        j.write_text(json.dumps(camp), encoding="utf-8")
        p = self.run_("depuis-json", j, "-o", s)
        self.assertEqual(p.returncode, 0, p.stdout)
        _, l = cl.annonces(cl.lire(s))
        self.assertEqual((l[0]["ad_group_id"], l[0]["status"], l[0]["image_url"], l[0].get("image_file_id", "")),
                         ("adgrp_101", "active", BON["image_url"], ""))

    def test_depuis_json_refuse_bloquant(self):
        j, s = self.d / "c.json", self.d / "s.xlsx"
        j.write_text(json.dumps({"ads": [dict(BON, title="x" * 40)]}), encoding="utf-8")
        p = self.run_("depuis-json", j, "-o", s)
        self.assertEqual(p.returncode, 3)
        self.assertFalse(s.exists())

    def test_tableau(self):
        f = self.d / "c.xlsx"
        classeur_brut(f, [BON])
        p = self.run_("tableau", f)
        self.assertEqual(p.stdout.splitlines()[0].split("\t"), cl.COLONNES)


class TestEnLigne(unittest.TestCase):
    PNG = b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\rIHDR" + (800).to_bytes(4, "big") + (800).to_bytes(4, "big")
    JPEG = (b"\xff\xd8\xff\xe0\x00\x10" + b"JFIF\x00" + b"\x00" * 9
            + b"\xff\xc0\x00\x11\x08" + (500).to_bytes(2, "big") + (700).to_bytes(2, "big") + b"\x03")
    WEBP = b"RIFF\x00\x00\x00\x00WEBPVP8X" + b"\x00" * 8 + (1199).to_bytes(3, "little") + (1199).to_bytes(3, "little")

    def test_dimensions(self):
        self.assertEqual(cl.dimensions_image(self.PNG), ("png", 800, 800))
        self.assertEqual(cl.dimensions_image(self.JPEG), ("jpeg", 700, 500))
        self.assertEqual(cl.dimensions_image(self.WEBP), ("webp", 1200, 1200))
        self.assertIsNone(cl.dimensions_image(b"<html>"))

    def test_controle(self):
        reponses = {BON["link"]: (200, "text/html", b""), "https://assurly.example/morte": (404, "", b""),
                    BON["image_url"]: (200, "image/png", self.PNG), "https://assurly.example/img/rect.jpg": (200, "image/jpeg", self.JPEG)}
        lignes = [dict(BON), dict(BON, link="https://assurly.example/morte", image_url="https://assurly.example/img/rect.jpg")]
        r = cl.Rapport()
        cl.controle_en_ligne(lignes, r, lire_url=lambda u: reponses[u])
        b = messages(r, cl.BLOQUANT)
        self.assertTrue(any("link répond 404" in m for m in b))
        self.assertTrue(any("carrée attendue" in m for m in b))
        self.assertTrue(any("640 × 640" in m for m in b))
        self.assertEqual(sum("annonce 1" == p["ou"] for p in r.points), 0)


if __name__ == "__main__":
    unittest.main()
